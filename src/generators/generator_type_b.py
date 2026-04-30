"""
Generator for Type-B reports – outputs an Excel workbook (.xlsx).

Replicates the source Hebrew RTL attendance sheet layout:
  Column order (A→L, displayed RTL):
    תאריך | יום | מקוד | כניסה | יציאה | הפסקה | סה"כ | 100% | 125% | 150% | שבת | הערות

  Summary box (bottom-right, same style as source):
    ימים, סה"כ שעות, 100% שעות, 125% שעות, 150% שעות, 150% שבת, בונוס, נסיעות

  Anomalous rows are highlighted in light-red; the "הערות" column (appended
  after שבת) contains the anomaly description.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.models import AttendanceReport
from src.generators.base_generator import BaseGenerator

# ── Palette ───────────────────────────────────────────────────────────────────
_FILL_HEADER  = PatternFill("solid", fgColor="4F4F4F")
_FILL_WHITE   = PatternFill("solid", fgColor="FFFFFF")
_FILL_ALT     = PatternFill("solid", fgColor="F2F2F2")
_FILL_WARN    = PatternFill("solid", fgColor="FFCCCC")
_FILL_TOTAL   = PatternFill("solid", fgColor="D9D9D9")
_FILL_SUM_LBL = PatternFill("solid", fgColor="D9D9D9")
_FILL_SUM_VAL = PatternFill("solid", fgColor="EFEFEF")

_FONT_HEADER  = Font(bold=True,  color="FFFFFF", name="Arial", size=9)
_FONT_BOLD    = Font(bold=True,  name="Arial", size=9)
_FONT_NORMAL  = Font(name="Arial", size=9)
_FONT_WARN    = Font(name="Arial", size=9, color="990000")
_FONT_TITLE   = Font(bold=True,  name="Arial", size=11)

_THIN    = Side(style="thin",   color="AAAAAA")
_MED     = Side(style="medium", color="888888")
_BORDER  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BORDER_MED = Border(left=_MED, right=_MED, top=_MED, bottom=_MED)

_AC = Alignment(horizontal="center", vertical="center")
_AR = Alignment(horizontal="right",  vertical="center")
_AL = Alignment(horizontal="left",   vertical="center")

# ── Columns: (header, width) – order matches source RTL sheet ─────────────────
_COLS = [
    ("תאריך",  13),
    ("יום",    11),
    ("מקוד",    7),
    ("כניסה",   8),
    ("יציאה",   8),
    ("הפסקה",   8),
    ('סה"כ',    8),
    ("100%",    7),
    ("125%",    7),
    ("150%",    7),
    ("שבת",     7),
    ("הערות",  42),
]
_NCOLS = len(_COLS)

_DAY_HE = {
    0: "יום שני",    1: "יום שלישי", 2: "יום רביעי",
    3: "יום חמישי", 4: "יום שישי",  5: "שבת",       6: "יום ראשון",
}

_SUMMARY_LABELS = [
    "ימים",
    'סה"כ שעות',
    "100% שעות",
    "125% שעות",
    "150% שעות",
    "150% שבת",
    "בונוס",
    "נסיעות",
]


def _c(ws, row, col, value="", font=None, fill=None, border=None, align=None):
    """Write a cell with optional formatting."""
    cell = ws.cell(row=row, column=col, value=value)
    if font:   cell.font      = font
    if fill:   cell.fill      = fill
    if border: cell.border    = border
    if align:  cell.alignment = align
    return cell


class GeneratorTypeB(BaseGenerator):
    def generate(self, report: AttendanceReport, output_path: Path) -> None:
        output_path = output_path.with_suffix(".xlsx")
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "דוח נוכחות"
        ws.sheet_view.rightToLeft = True

        # ── Row 1: company name ───────────────────────────────────────────────
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_NCOLS)
        _c(ws, 1, 1, 'נ.ע. הנסר כח אדם בע"מ',
           font=_FONT_TITLE, align=_AR)
        ws.row_dimensions[1].height = 22

        # ── Row 2: employee + period ──────────────────────────────────────────
        emp_line = (
            f"עובד: {report.employee or '—'}"
            f"    תקופה: {report.period_start} – {report.period_end}"
        )
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=_NCOLS)
        _c(ws, 2, 1, emp_line, font=_FONT_BOLD, align=_AR)
        ws.row_dimensions[2].height = 16

        # blank row 3
        ws.row_dimensions[3].height = 6

        # ── Row 4: column headers ─────────────────────────────────────────────
        HDR_ROW = 4
        for ci, (hdr, _) in enumerate(_COLS, 1):
            _c(ws, HDR_ROW, ci, hdr,
               font=_FONT_HEADER, fill=_FILL_HEADER, border=_BORDER, align=_AC)
        ws.row_dimensions[HDR_ROW].height = 18

        # ── Data rows ─────────────────────────────────────────────────────────
        n = len(report.records)
        for idx, rec in enumerate(report.records):
            ri     = HDR_ROW + 1 + idx
            is_bad = bool(rec.anomalies)
            fill   = _FILL_WARN if is_bad else (_FILL_ALT if idx % 2 else _FILL_WHITE)
            font   = _FONT_WARN if is_bad else _FONT_NORMAL

            day_he = ""
            for fmt in ("%d/%m/%Y", "%d/%m/%y"):
                try:
                    day_he = _DAY_HE.get(
                        datetime.strptime(rec.date_text, fmt).weekday(), "")
                    break
                except ValueError:
                    pass

            values = [
                rec.date_text,           # תאריך
                day_he,                  # יום
                "",                      # מקוד – not parsed
                rec.start_time,          # כניסה
                rec.end_time,            # יציאה
                rec.break_hours or "",   # הפסקה
                rec.daily_hours,         # סה"כ
                "", "", "", "",          # 100% 125% 150% שבת
                "; ".join(rec.anomalies),# הערות
            ]
            for ci, val in enumerate(values, 1):
                _c(ws, ri, ci, val, font=font, fill=fill, border=_BORDER, align=_AC)
            ws.row_dimensions[ri].height = 16

        # ── Total row ─────────────────────────────────────────────────────────
        total_row = HDR_ROW + 1 + n
        for ci in range(1, _NCOLS + 1):
            _c(ws, total_row, ci, "", font=_FONT_BOLD, fill=_FILL_TOTAL,
               border=_BORDER, align=_AC)
        ws.cell(total_row, 1).value     = str(n)           # day count
        ws.cell(total_row, 7).value     = report.total_hours  # סה"כ col
        ws.cell(total_row, 7).font      = _FONT_BOLD
        ws.row_dimensions[total_row].height = 18

        # ── Summary box ───────────────────────────────────────────────────────
        total_h      = report.total_hours
        summary_vals = [n, total_h, total_h, 0, 0, 0, 0, 0]
        sr = total_row + 2
        for i, (label, val) in enumerate(zip(_SUMMARY_LABELS, summary_vals)):
            r = sr + i
            _c(ws, r, 1, label, font=_FONT_BOLD,   fill=_FILL_SUM_LBL,
               border=_BORDER, align=_AR)
            _c(ws, r, 2, val,   font=_FONT_NORMAL, fill=_FILL_SUM_VAL,
               border=_BORDER, align=_AC)
            ws.row_dimensions[r].height = 16

        # ── Column widths ──────────────────────────────────────────────────────
        for ci, (_, width) in enumerate(_COLS, 1):
            ws.column_dimensions[get_column_letter(ci)].width = width

        # ── Freeze below header ───────────────────────────────────────────────
        ws.freeze_panes = ws.cell(HDR_ROW + 1, 1)

        wb.save(str(output_path))
